#!/usr/bin/env python3
"""
第1章教材生成器 V4 — 语法高亮 + PDF + DOCX
"""
import json, re, os, sys
from pathlib import Path

BOOK_ROOT = Path(__file__).parent.parent

# ============================================================
# 14道题数据 (同上)
# ============================================================
PROBLEMS = [
    {"nq": "NQ001", "acw": 1, "title": "A + B",
     "desc": "小鲁第一次打开TRAE，AI助手提示他：\"想学编程，从最简单的计算开始。\"输入两个整数A和B，计算它们的和。",
     "input_fmt": "一行，两个整数A和B，用空格隔开。",
     "output_fmt": "一个整数，即A+B的结果。",
     "sample_in": "3 4", "sample_out": "7",
     "constraint": "0 ≤ A, B ≤ 10⁸",
     "solution": "这是编程中最简单的题目，但它包含了所有程序的基本骨架：读入数据→计算→输出结果。C++需要引入iostream、声明main函数、用cin/cout；Python更简洁。两种语言体现了静态类型与动态类型的设计哲学差异。",
     "technique": "Python的map(int, input().split())是处理空格分隔整数输入的标准写法。C++的cin >> a >> b自动跳过空白字符。",
     "cpp": '#include <iostream>\nusing namespace std;\nint main() {\n    int a, b;\n    cin >> a >> b;\n    cout << a + b << endl;\n    return 0;\n}',
     "py": 'a, b = map(int, input().split())\nprint(a + b)'},
    {"nq": "NQ002", "acw": 608, "title": "差",
     "desc": "老师给出四个整数A、B、C、D，请计算A×B−C×D的结果。",
     "input_fmt": "一行，四个整数A、B、C、D，用空格隔开。",
     "output_fmt": "输出\"DIFERENCA = \"后跟计算结果。",
     "sample_in": "5 6 7 8", "sample_out": "DIFERENCA = -26",
     "constraint": "−10⁴ ≤ A,B,C,D ≤ 10⁴",
     "solution": "直接按公式计算，注意运算顺序：先乘后减。C++用int可满足范围，Python的int无范围限制。输出需要带前缀\"DIFERENCA = \"。",
     "technique": "Python的f-string：f\"DIFERENCA = {a*b - c*d}\"，比%和.format()更清晰直观。",
     "cpp": '#include <cstdio>\nint main() {\n    int a, b, c, d;\n    scanf("%d%d%d%d", &a, &b, &c, &d);\n    printf("DIFERENCA = %d\\n", a * b - c * d);\n    return 0;\n}',
     "py": 'a, b, c, d = map(int, input().split())\nprint(f"DIFERENCA = {a * b - c * d}")'},
    {"nq": "NQ003", "acw": 604, "title": "圆的面积",
     "desc": "小鲁在几何课上学到了圆的面积公式S=πr²。给定半径r，计算圆的面积。π取3.14159。",
     "input_fmt": "一个浮点数r，表示圆的半径。",
     "output_fmt": "输出\"A=\"后跟面积，保留4位小数。",
     "sample_in": "2.00", "sample_out": "A=12.5664",
     "constraint": "0 < r ≤ 10000",
     "solution": "面积 = π × r²。使用浮点类型：C++的double、Python的float。C++用printf控制小数位；Python用f-string。",
     "technique": "C++格式化输出：printf(\"%.4lf\", x)简洁；cout<<fixed<<setprecision(4)<<x类型安全。Python的f\"{x:.4f}\"最直观。",
     "cpp": '#include <cstdio>\nint main() {\n    double r;\n    scanf("%lf", &r);\n    printf("A=%.4lf\\n", 3.14159 * r * r);\n    return 0;\n}',
     "py": 'r = float(input())\nprint(f"A={3.14159 * r * r:.4f}")'},
    {"nq": "NQ004", "acw": 606, "title": "平均数1",
     "desc": "学期结束，老师需要计算小鲁的加权平均成绩。A科目的权重是3.5，B科目的权重是7.5。",
     "input_fmt": "两行，每行一个浮点数（成绩A和B）。",
     "output_fmt": "\"MEDIA = \"后跟加权平均分，保留5位小数。",
     "sample_in": "5.0\n7.1", "sample_out": "MEDIA = 6.43182",
     "constraint": "0 ≤ A, B ≤ 10.0",
     "solution": "加权平均公式：(A×3.5 + B×7.5) / 11。分母是权重之和(3.5+7.5=11)，不是题目数2。这是常见陷阱。",
     "technique": "C++所有浮点字面量默认double。Python的/总是返回浮点数，无需担心整数除法问题。",
     "cpp": '#include <cstdio>\nint main() {\n    double a, b;\n    scanf("%lf%lf", &a, &b);\n    printf("MEDIA = %.5lf\\n", (a * 3.5 + b * 7.5) / 11);\n    return 0;\n}',
     "py": 'a = float(input())\nb = float(input())\nprint(f"MEDIA = {(a * 3.5 + b * 7.5) / 11:.5f}")'},
    {"nq": "NQ005", "acw": 609, "title": "工资",
     "desc": "小鲁暑假打工，公司需要计算他的工资。已知员工编号、月工作小时数和时薪，请计算月工资总额。",
     "input_fmt": "三行：编号（整数）、时数（整数）、时薪（浮点数）。",
     "output_fmt": "\"NUMBER = X\"换行\"SALARY = U$ XX.XX\"。",
     "sample_in": "25\n100\n5.50", "sample_out": "NUMBER = 25\nSALARY = U$ 550.00",
     "constraint": "1 ≤ 编号 ≤ 100, 1 ≤ 时数 ≤ 200, 1 ≤ 时薪 ≤ 50",
     "solution": "工资 = 时数 × 时薪。输入包含整数和浮点数混合类型，C++用scanf精确控制格式。",
     "technique": "C++中scanf可精确控制混合输入。Python逐行读取更清晰安全。",
     "cpp": '#include <cstdio>\nint main() {\n    int n, h;\n    double m;\n    scanf("%d%d%lf", &n, &h, &m);\n    printf("NUMBER = %d\\nSALARY = U$ %.2lf\\n", n, h * m);\n    return 0;\n}',
     "py": 'n = int(input())\nh = int(input())\nm = float(input())\nprint(f"NUMBER = {n}")\nprint(f"SALARY = U$ {h * m:.2f}")'},
    {"nq": "NQ006", "acw": 615, "title": "油耗",
     "desc": "小鲁买了一辆二手车，想测试油耗。记录行驶总距离（公里）和消耗汽油量（升），计算每升汽油可行驶公里数。",
     "input_fmt": "两行：距离（浮点数, km）、汽油量（浮点数, L）。",
     "output_fmt": "每升公里数，保留3位小数，后跟\" km/l\"。",
     "sample_in": "500\n35.0", "sample_out": "14.286 km/l",
     "constraint": "1 ≤ 距离 ≤ 10⁶, 0 < 汽油量 ≤ 10⁵",
     "solution": "燃油效率 = 距离 / 汽油量。使用浮点数除法，输出保留3位小数。",
     "technique": "C++整数除法截断，距离如果是int需转double。Python中/自动返回浮点数。",
     "cpp": '#include <cstdio>\nint main() {\n    double x, y;\n    scanf("%lf%lf", &x, &y);\n    printf("%.3lf km/l\\n", x / y);\n    return 0;\n}',
     "py": 'x = float(input())\ny = float(input())\nprint(f"{x / y:.3f} km/l")'},
    {"nq": "NQ007", "acw": 616, "title": "两点间的距离",
     "desc": "小鲁在坐标系上标了两个点P1(x1,y1)和P2(x2,y2)，他想知道这两点的欧几里得距离。",
     "input_fmt": "一行，四个浮点数x1, y1, x2, y2。",
     "output_fmt": "两点间距离，保留4位小数。",
     "sample_in": "1.0 7.0 5.0 9.0", "sample_out": "4.4721",
     "constraint": "−10⁹ ≤ 坐标 ≤ 10⁹",
     "solution": "距离 = √((x1−x2)² + (y1−y2)²)。第一次用到数学库。注意使用double确保精度。",
     "technique": "C++需#include <cmath>。Python的math.sqrt()简洁明了。f\"{dist:.4f}\"格式化。",
     "cpp": '#include <cstdio>\n#include <cmath>\nint main() {\n    double x1,y1,x2,y2;\n    scanf("%lf%lf%lf%lf",&x1,&y1,&x2,&y2);\n    double d=sqrt((x1-x2)*(x1-x2)+(y1-y2)*(y1-y2));\n    printf("%.4lf\\n",d);\n    return 0;\n}',
     "py": 'import math\nx1,y1,x2,y2=map(float,input().split())\nd=math.sqrt((x1-x2)**2+(y1-y2)**2)\nprint(f"{d:.4f}")'},
    {"nq": "NQ008", "acw": 653, "title": "钞票",
     "desc": "小鲁在银行取钱，ATM机需要给出最少数量的钞票。给定金额N，计算需要多少张各面额钞票。",
     "input_fmt": "一个整数N。",
     "output_fmt": "第一行输出N，然后7行按面额从大到小输出张数。",
     "sample_in": "576", "sample_out": "576\n5 nota(s) de R$ 100,00\n1 nota(s) de R$ 50,00\n1 nota(s) de R$ 20,00\n0 nota(s) de R$ 10,00\n1 nota(s) de R$ 5,00\n0 nota(s) de R$ 2,00\n1 nota(s) de R$ 1,00",
     "constraint": "0 < N < 10⁶",
     "solution": "贪心算法雏形：count = N / face_value; N %= face_value。Python用列表+循环消除重复代码——本课最重要的工程思维。",
     "technique": "Python用列表+循环消除重复：for v in [100,50,20,10,5,2,1]。对比C++的7段重复代码。",
     "cpp": '#include <cstdio>\nint main() {\n    int n, v[]={100,50,20,10,5,2,1};\n    scanf("%d",&n); printf("%d\\n",n);\n    for(int i=0;i<7;i++){\n        printf("%d nota(s) de R$ %d,00\\n",n/v[i],v[i]);\n        n%=v[i];\n    }\n    return 0;\n}',
     "py": 'n=int(input())\nprint(n)\nfor v in[100,50,20,10,5,2,1]:\n    print(f"{n//v} nota(s) de R$ {v},00")\n    n%=v'},
    {"nq": "NQ009", "acw": 654, "title": "时间转换",
     "desc": "小鲁的计时器只显示秒数，他想转换为\"时:分:秒\"格式。给定总秒数N，转换成HH:MM:SS。",
     "input_fmt": "一个整数N。",
     "output_fmt": "HH:MM:SS格式。",
     "sample_in": "556", "sample_out": "0:9:16",
     "constraint": "0 ≤ N ≤ 10⁶",
     "solution": "时=N/3600, 分=N%3600/60, 秒=N%60。Python的divmod()一次性获得商和余数。",
     "technique": "Python的divmod(a,b)返回(商,余数)元组，优雅解包接收。C++需分开计算/和%。",
     "cpp": '#include <cstdio>\nint main() {\n    int n;\n    scanf("%d",&n);\n    printf("%d:%d:%d",n/3600,n%3600/60,n%60);\n    return 0;\n}',
     "py": 'n=int(input())\nh,r=divmod(n,3600)\nm,s=divmod(r,60)\nprint(f"{h}:{m}:{s}")'},
    {"nq": "NQ010", "acw": 605, "title": "简单乘积",
     "desc": "小鲁发现乘法比加法快得多。给定两个整数，计算它们的乘积。",
     "input_fmt": "两行，每行一个整数。",
     "output_fmt": "\"PROD = \"后跟乘积。",
     "sample_in": "3\n9", "sample_out": "PROD = 27",
     "constraint": "−10⁴ ≤ A, B ≤ 10⁴",
     "solution": "与NQ001结构相同，将+改成*并添加\"PROD = \"前缀。让学生独立完成一次完整的修改→编译→AC流程。",
     "technique": "修改已有代码比从零开始更高效。用TRAE打开NQ001的代码，让AI改成乘法版本。",
     "cpp": '#include <iostream>\nusing namespace std;\nint main(){\n    int a,b;cin>>a>>b;\n    cout<<"PROD = "<<a*b<<endl;\n    return 0;\n}',
     "py": 'a=int(input())\nb=int(input())\nprint(f"PROD = {a*b}")'},
    {"nq": "NQ011", "acw": 611, "title": "简单计算",
     "desc": "根据产品编号、数量和单价，计算总价。第一行两个整数，第二行一个浮点数。",
     "input_fmt": "第一行：code和quantity（整数）。第二行：price（浮点数）。",
     "output_fmt": "\"VALOR A PAGAR: R$ XX.XX\"。",
     "sample_in": "12 1\n5.30", "sample_out": "VALOR A PAGAR: R$ 5.30",
     "constraint": "1 ≤ code ≤ 100, 1 ≤ quantity ≤ 100",
     "solution": "总价 = 数量 × 单价。int × float自动提升类型。",
     "technique": "C++中int×double自动转double。Python中int×float自动转float。",
     "cpp": '#include <cstdio>\nint main(){\n    int c,q;double p;\n    scanf("%d%d%lf",&c,&q,&p);\n    printf("VALOR A PAGAR: R$ %.2lf\\n",q*p);\n    return 0;\n}',
     "py": 'c,q=map(int,input().split())\np=float(input())\nprint(f"VALOR A PAGAR: R$ {q*p:.2f}")'},
    {"nq": "NQ012", "acw": 612, "title": "球的体积",
     "desc": "小鲁在物理课上学了球体体积公式V=(4/3)πr³。给定半径r，计算体积。π=3.14159。",
     "input_fmt": "一个整数r。",
     "output_fmt": "\"VOLUME = \"后跟体积，保留3位小数。",
     "sample_in": "3", "sample_out": "VOLUME = 113.097",
     "constraint": "1 ≤ r ≤ 1000",
     "solution": "V=(4.0/3.0)×π×r³。C++整数除法陷阱：4/3=1，必须写4.0/3.0。",
     "technique": "C++中整数除法4/3=1的陷阱是初学者最容易犯的错误。Python自动浮点。",
     "cpp": '#include <cstdio>\nint main(){\n    int r;scanf("%d",&r);\n    printf("VOLUME = %.3lf\\n",4.0/3.0*3.14159*r*r*r);\n    return 0;\n}',
     "py": 'r=int(input())\nv=4.0/3.0*3.14159*r**3\nprint(f"VOLUME = {v:.3f}")'},
    {"nq": "NQ013", "acw": 613, "title": "面积",
     "desc": "计算三个几何图形的面积：(1)直角三角形A*C/2；(2)圆π*C²；(3)梯形(A+B)*C/2。π=3.14159。",
     "input_fmt": "一行，三个浮点数A、B、C。",
     "output_fmt": "三行：TRIANGULO/CIRCULO/TRAPEZIO，各保留3位小数。",
     "sample_in": "3.0 4.0 5.2", "sample_out": "TRIANGULO: 7.800\nCIRCULO: 84.949\nTRAPEZIO: 18.200",
     "constraint": "0 < A, B, C ≤ 100",
     "solution": "三道公式一题，注意A、B、C在不同公式中担任不同角色。使用常量PI。",
     "technique": "一道题含三个独立计算，是测试代码组织能力的好题。",
     "cpp": '#include <cstdio>\nint main(){\n    double a,b,c;\n    scanf("%lf%lf%lf",&a,&b,&c);\n    printf("TRIANGULO: %.3lf\\n",a*c/2);\n    printf("CIRCULO: %.3lf\\n",3.14159*c*c);\n    printf("TRAPEZIO: %.3lf\\n",(a+b)*c/2);\n    return 0;\n}',
     "py": 'a,b,c=map(float,input().split())\nprint(f"TRIANGULO: {a*c/2:.3f}")\nprint(f"CIRCULO: {3.14159*c*c:.3f}")\nprint(f"TRAPEZIO: {(a+b)*c/2:.3f}")'},
    {"nq": "NQ014", "acw": 614, "title": "最大值",
     "desc": "输入三个整数a、b、c，输出其中的最大者。这是学习比较和选择逻辑的收尾题。",
     "input_fmt": "一行，三个整数。",
     "output_fmt": "一个整数，即最大值。",
     "sample_in": "7 14 106", "sample_out": "106",
     "constraint": "−10⁹ ≤ a, b, c ≤ 10⁹",
     "solution": "Python的max(a,b,c)直接接收多参数。C++需max({a,b,c})或max(a,max(b,c))。体现了Python\"电池已包含\"的设计哲学。",
     "technique": "Python内置max/min/sum/abs等函数。C++的<algorithm>提供类似功能但语法不如Python自然。",
     "cpp": '#include <iostream>\n#include <algorithm>\nusing namespace std;\nint main(){\n    int a,b,c;cin>>a>>b>>c;\n    cout<<max({a,b,c})<<endl;\n    return 0;\n}',
     "py": 'a,b,c=map(int,input().split())\nprint(max(a,b,c))'},
]

# ============================================================
# 语法高亮 (pygments)
# ============================================================
from pygments import highlight
from pygments.lexers import CppLexer, PythonLexer
from pygments.formatters import HtmlFormatter

# 'vs' 主题：白底，鲜明配色，适合印刷
CPP_FORMATTER = HtmlFormatter(style='vs', noclasses=True)
PY_FORMATTER = HtmlFormatter(style='vs', noclasses=True)

def highlight_code(code, lexer, formatter):
    """用 pygments 高亮代码，返回带行内样式的 <pre> 内容（保留换行和缩进）"""
    html = highlight(code, lexer, formatter)
    # 提取 <pre> 内的内容（保留 \n 作为真实换行）
    m = re.search(r'<pre[^>]*>(.*)</pre>', html, re.DOTALL)
    if m:
        inner = m.group(1)
        # pygments 把每行包装在 <span> 里，行间是真实换行符
        # 去掉外层空 span
        inner = re.sub(r'<span></span>\n?', '', inner)
        return inner
    return code

def highlight_cpp(code):
    return highlight_code(code, CppLexer(), CPP_FORMATTER)

def highlight_py(code):
    return highlight_code(code, PythonLexer(), PY_FORMATTER)

# ============================================================
# CSS — 书本排版
# ============================================================
CSS = r"""
@page {
    size: A4;
    margin: 2.2cm 2cm 2.2cm 2cm;
    @top-center { content: string(chapter); font-size: 7.5pt; color: #999; font-family: "PingFang SC", sans-serif; }
    @bottom-center { content: counter(page); font-size: 7.5pt; color: #999; }
}

body {
    font-family: "PingFang SC", "Hiragino Sans GB", "Noto Serif CJK SC", "STSong", serif;
    font-size: 9.5pt;
    line-height: 1.7;
    color: #222;
    string-set: chapter "第1章 程序设计的第一个脚印 —— 变量、输入输出与顺序结构";
}

.chapter-title { text-align: center; font-size: 20pt; font-weight: bold; margin: 1.5em 0 0.1em 0; letter-spacing: 3pt; }
.chapter-subtitle { text-align: center; font-size: 10pt; color: #777; margin-bottom: 1.5em; padding-bottom: 0.8em; border-bottom: 1px solid #bbb; }

.preface { font-size: 10pt; margin-bottom: 2em; color: #444; }
.preface p { margin: 0.4em 0; text-indent: 2em; }

.problem-title { font-size: 12pt; font-weight: bold; margin: 1.5em 0 0.4em 0; padding-bottom: 0.15em; border-bottom: 1pt solid #444; }
.problem-title .nq { color: #2563eb; margin-right: 0.6em; font-size: 11pt; }
.problem-title .acw { font-size: 7.5pt; color: #aaa; font-weight: normal; margin-left: 1em; }

.problem-desc { margin: 0.6em 0 1em 0; text-indent: 2em; line-height: 1.85; }

.spec-table { width: 100%; border-collapse: collapse; margin: 0.3em 0 0.5em 0; font-size: 9pt; }
.spec-table td { padding: 0.3em 0.8em; vertical-align: top; border: none; }
.spec-table .spec-label { width: 4em; font-size: 8pt; font-weight: bold; text-align: right; padding-right: 1em; white-space: nowrap; }
.spec-table .spec-label .tag { display: inline-block; padding: 0.15em 0.5em; border-radius: 2px; color: #fff; font-size: 7.5pt; letter-spacing: 0.5pt; }
.spec-table .spec-label .tag.in  { background: #2563eb; }
.spec-table .spec-label .tag.out { background: #059669; }
.spec-table .spec-label .tag.lim { background: #d97706; }
.spec-table .spec-value { color: #333; font-size: 9pt; }

.sample-box { background: #f7f8fa; border: 0.5pt solid #dde; border-radius: 4px; padding: 0.7em 1em; margin: 1em 0 1.2em 0; }
.sample-grid { display: flex; gap: 2em; }
.sample-col { flex: 1; }
.sample-col .col-label { font-size: 7.5pt; color: #888; margin-bottom: 0.2em; font-weight: bold; }
.sample-col pre { background: none; border: none; padding: 0.3em 0; margin: 0; font-family: "SF Mono","Menlo","Consolas",monospace; font-size: 9pt; line-height: 1.4; white-space: pre-wrap; color: #333; }

.insight-block { margin: 0.8em 0; padding: 0.5em 0.8em; border-left: 3px solid #2563eb; background: #f8faff; }
.insight-block .insight-label { font-size: 8pt; font-weight: bold; color: #2563eb; margin-right: 0.5em; }

/* 双栏代码 — 语法高亮 */
.code-dual { display: flex; gap: 1.2em; margin: 1.2em 0; page-break-inside: avoid; }
.code-col { min-width: 0; }
.code-col:first-child { flex: 3; }
.code-col:last-child { flex: 2; }
.code-col .lang-badge { display: inline-block; font-size: 7.5pt; font-weight: bold; color: #fff; background: #2563eb; padding: 0.2em 0.7em; border-radius: 3px; margin-bottom: 0.4em; }
.code-col pre { background: #f8f8f0; border: 0.5pt solid #e0e0e0; border-radius: 4px; padding: 0.7em 0.9em; font-family: "SF Mono","Menlo","Consolas","Courier New",monospace; font-size: 7.5pt; line-height: 1.45; overflow-x: auto; margin: 0; white-space: pre-wrap; word-break: break-all; }

.section-divider { border: none; border-top: 0.3pt solid #e0e0e0; margin: 1em 0 0 0; }
.chapter-end { text-align: center; margin-top: 3em; font-size: 8pt; color: #999; }
"""

def build_problem_html(p):
    parts = ['<div class="problem-block">']
    parts.append(f'<div class="problem-title"><span class="nq">{p["nq"]}</span>{p["title"]}<span class="acw">AcWing {p["acw"]}</span></div>')
    parts.append(f'<div class="problem-desc">{p["desc"]}</div>')
    parts.append('<table class="spec-table">')
    parts.append(f'<tr><td class="spec-label"><span class="tag in">输入</span></td><td class="spec-value">{p["input_fmt"]}</td></tr>')
    parts.append(f'<tr><td class="spec-label"><span class="tag out">输出</span></td><td class="spec-value">{p["output_fmt"]}</td></tr>')
    parts.append(f'<tr><td class="spec-label"><span class="tag lim">范围</span></td><td class="spec-value">{p["constraint"]}</td></tr>')
    parts.append('</table>')
    parts.append('<div class="sample-box"><div class="sample-grid">')
    parts.append(f'<div class="sample-col"><div class="col-label">输入</div><pre>{p["sample_in"]}</pre></div>')
    parts.append(f'<div class="sample-col"><div class="col-label">输出</div><pre>{p["sample_out"]}</pre></div>')
    parts.append('</div></div>')
    parts.append(f'<div class="insight-block"><span class="insight-label">思路</span><span>{p["solution"]}</span></div>')
    parts.append(f'<div class="insight-block"><span class="insight-label">技巧</span><span>{p["technique"]}</span></div>')

    # 语法高亮代码
    cpp_hl = highlight_cpp(p["cpp"])
    py_hl = highlight_py(p["py"])
    parts.append('<div class="code-dual">')
    parts.append(f'<div class="code-col"><div class="lang-badge">C++</div><pre>{cpp_hl}</pre></div>')
    parts.append(f'<div class="code-col"><div class="lang-badge">Python</div><pre>{py_hl}</pre></div>')
    parts.append('</div>')
    parts.append('</div><hr class="section-divider">')
    return '\n'.join(parts)

def generate_html():
    problems_html = [build_problem_html(p) for p in PROBLEMS]
    html = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">
<title>第1章 程序设计的第一个脚印</title><style>{CSS}</style></head><body>
<div class="chapter-title">程序设计的第一个脚印</div>
<div class="chapter-subtitle">变量、输入输出与顺序结构 · 14题 · C++ &amp; Python 双语对照</div>
<div class="preface">
<p>本章是编程之旅的第一站。14道题目从最简单的A+B开始，逐步引入浮点数运算、格式化输出、数学库、整除取模、时间换算和几何公式。每道题都同时提供C++和Python两种实现，读者可以对照阅读，体会两种语言的设计哲学差异。</p>
<p>如果你正在使用TRAE这类AI编程工具学习，建议先自己思考解题思路，然后对照书中的代码理解每一行的含义。记住：<strong>读懂代码比写出代码更重要</strong>。</p>
</div>
{chr(10).join(problems_html)}
<div class="chapter-end">— 第1章完 · 共14题 —</div>
</body></html>"""
    out = BOOK_ROOT / "textbook" / "chapter01_print.html"
    out.write_text(html, encoding="utf-8")
    return out


# ============================================================
# DOCX 生成
# ============================================================
def generate_docx():
    """生成与PDF排版一致的DOCX文档"""
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml

    doc = Document()

    # 页面设置 A4
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # 默认样式
    style = doc.styles['Normal']
    style.font.name = '等线'
    style.font.size = Pt(9.5)
    style.paragraph_format.line_spacing = 1.35
    style.paragraph_format.space_after = Pt(3)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '等线')

    BLUE = RGBColor(0x25, 0x63, 0xEB)
    GREEN = RGBColor(0x05, 0x96, 0x69)
    ORANGE = RGBColor(0xD9, 0x77, 0x06)
    GRAY = RGBColor(0x66, 0x66, 0x66)
    LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
    CODE_BG = 'F0F0F0'
    SAMPLE_BG = 'F7F8FA'

    def add_colored_tag(para, text, color, bold=True):
        run = para.add_run(text)
        run.font.size = Pt(8)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # use paragraph shading for tag background
        return run

    def add_shaded_paragraph(doc, text, bg_color, font_size=Pt(9)):
        """添加带背景色的段落"""
        para = doc.add_paragraph()
        pPr = para._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        pPr.append(shd)
        run = para.add_run(text)
        run.font.size = font_size
        return para

    def add_code_run(para, text, color=None, bold=False, italic=False, font_size=Pt(7.5)):
        """添加一个代码片段run"""
        run = para.add_run(text)
        run.font.name = 'Courier New'
        run.font.size = font_size
        if color:
            run.font.color.rgb = color
        if bold:
            run.font.bold = True
        if italic:
            run.font.italic = True
        return run

    # ===== 标题 =====
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run('程序设计的第一个脚印')
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_para.paragraph_format.space_after = Pt(2)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run('变量、输入输出与顺序结构 · 14题 · C++ & Python 双语对照')
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = GRAY
    sub_para.paragraph_format.space_after = Pt(8)

    # 分隔线（用border模拟）
    border_para = doc.add_paragraph()
    pPr = border_para._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="4" w:space="1" w:color="BBBBBB"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    border_para.paragraph_format.space_after = Pt(12)

    # ===== 前言 =====
    preface = doc.add_paragraph(
        '本章是编程之旅的第一站。14道题目从最简单的A+B开始，逐步引入浮点数运算、格式化输出、'
        '数学库、整除取模、时间换算和几何公式。每道题都同时提供C++和Python两种实现，'
        '读者可以对照阅读，体会两种语言的设计哲学差异。'
    )
    preface.paragraph_format.first_line_indent = Cm(0.7)
    preface.paragraph_format.space_after = Pt(6)

    preface2 = doc.add_paragraph(
        '如果你正在使用TRAE这类AI编程工具学习，建议先自己思考解题思路，然后对照书中的代码'
        '理解每一行的含义。记住：读懂代码比写出代码更重要。'
    )
    preface2.paragraph_format.first_line_indent = Cm(0.7)
    preface2.paragraph_format.space_after = Pt(12)

    # ===== 每题 =====
    from pygments.lexers import CppLexer, PythonLexer
    from pygments.token import Token

    for p in PROBLEMS:
        # 标题
        title_para = doc.add_paragraph()
        title_para.paragraph_format.space_before = Pt(14)
        title_para.paragraph_format.space_after = Pt(4)
        nq_run = title_para.add_run(f'{p["nq"]} ')
        nq_run.font.size = Pt(12)
        nq_run.font.bold = True
        nq_run.font.color.rgb = BLUE
        t_run = title_para.add_run(p['title'])
        t_run.font.size = Pt(12)
        t_run.font.bold = True
        acw_run = title_para.add_run(f'  AcWing {p["acw"]}')
        acw_run.font.size = Pt(7.5)
        acw_run.font.color.rgb = LIGHT_GRAY

        # 标题下划线
        pPr = title_para._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="444444"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

        # 题面描述
        desc_para = doc.add_paragraph(p['desc'])
        desc_para.paragraph_format.first_line_indent = Cm(0.7)
        desc_para.paragraph_format.space_after = Pt(4)

        # 规格表（输入/输出/范围） — 用 docx 表格
        spec_table = doc.add_table(rows=3, cols=2)
        spec_table.autofit = True
        spec_table.alignment = WD_TABLE_ALIGNMENT.LEFT

        spec_data = [
            ('输入', p['input_fmt'], BLUE),
            ('输出', p['output_fmt'], GREEN),
            ('范围', p['constraint'], ORANGE),
        ]
        for i, (label, value, color) in enumerate(spec_data):
            cell_label = spec_table.cell(i, 0)
            cell_label.width = Cm(1.5)
            cell_label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            # 标签
            tag_para = cell_label.paragraphs[0]
            tag_para.paragraph_format.space_after = Pt(0)
            tag_run = tag_para.add_run(f' {label} ')
            tag_run.font.size = Pt(7.5)
            tag_run.font.bold = True
            tag_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            # 标签背景色
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>')
            tag_run._r.get_or_add_rPr().append(shd)

            cell_val = spec_table.cell(i, 1)
            cell_val.width = Cm(14)
            val_para = cell_val.paragraphs[0]
            val_para.paragraph_format.space_after = Pt(0)
            val_run = val_para.add_run(value)
            val_run.font.size = Pt(9)

        # 样例 — 用表格：左栏输入，右栏输出
        sample_table = doc.add_table(rows=1, cols=2)
        sample_table.autofit = True
        sample_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        # 浅灰底色
        for col_idx in [0, 1]:
            cell = sample_table.cell(0, col_idx)
            cell.width = Cm(7.5)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{SAMPLE_BG}" w:val="clear"/>')
            tcPr.append(shd)
            # 标签行
            label_para = cell.paragraphs[0]
            label_para.paragraph_format.space_after = Pt(2)
            ln = '输入' if col_idx == 0 else '输出'
            lr = label_para.add_run(ln)
            lr.font.size = Pt(7.5)
            lr.font.bold = True
            lr.font.color.rgb = LIGHT_GRAY
            # 内容行
            val_para = cell.add_paragraph()
            val_para.paragraph_format.space_after = Pt(2)
            vr = val_para.add_run(p['sample_in'] if col_idx == 0 else p['sample_out'])
            vr.font.name = 'Courier New'
            vr.font.size = Pt(9)

        # 思路
        insight1 = doc.add_paragraph()
        insight1.paragraph_format.space_after = Pt(4); insight1.paragraph_format.space_before = Pt(4)
        pPr = insight1._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFF" w:val="clear"/>')
        pPr.append(shd)
        lbl1 = insight1.add_run('思路 ')
        lbl1.font.size = Pt(8)
        lbl1.font.bold = True
        lbl1.font.color.rgb = BLUE
        t1 = insight1.add_run(p['solution'])
        t1.font.size = Pt(9)

        # 技巧
        insight2 = doc.add_paragraph()
        insight2.paragraph_format.space_after = Pt(6); insight2.paragraph_format.space_before = Pt(2)
        pPr = insight2._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFF" w:val="clear"/>')
        pPr.append(shd)
        lbl2 = insight2.add_run('技巧 ')
        lbl2.font.size = Pt(8)
        lbl2.font.bold = True
        lbl2.font.color.rgb = BLUE
        t2 = insight2.add_run(p['technique'])
        t2.font.size = Pt(9)

        # 双栏代码 — 两个表格列，等宽
        code_table = doc.add_table(rows=1, cols=2)
        code_table.autofit = False
        code_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for col_idx, (lang, code, lexer) in enumerate([
            ('C++', p['cpp'], CppLexer()),
            ('Python', p['py'], PythonLexer()),
        ]):
            cell = code_table.cell(0, col_idx)
            # 每列约 7.5cm
            cell.width = Cm(7.5)
            cell.paragraphs[0].clear()

            # 语言标签
            badge_para = cell.paragraphs[0]
            badge_para.paragraph_format.space_after = Pt(2)
            badge_run = badge_para.add_run(f' {lang} ')
            badge_run.font.size = Pt(7)
            badge_run.font.bold = True
            badge_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="2563EB" w:val="clear"/>')
            badge_run._r.get_or_add_rPr().append(shd)

            # 代码（语法高亮）— 每行一个 paragraph
            lines = code.split('\n')
            for line_idx, line in enumerate(lines):
                code_para = cell.add_paragraph()
                code_para.paragraph_format.space_before = Pt(0)
                code_para.paragraph_format.space_after = Pt(0)
                code_para.paragraph_format.line_spacing = 1.1
                # 浅灰背景
                pPr = code_para._p.get_or_add_pPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CODE_BG}" w:val="clear"/>')
                pPr.append(shd)

                if not line.strip():
                    # 空行：加个空格保持行高
                    add_code_run(code_para, ' ', font_size=Pt(6))
                    continue

                # pygments tokenize
                tokens = list(lexer.get_tokens(line))
                for token_type, token_text in tokens:
                    if not token_text:
                        continue
                    color_map = {
                        Token.Keyword: RGBColor(0x00, 0x00, 0xFF),
                        Token.Keyword.Type: RGBColor(0x2B, 0x91, 0xAF),
                        Token.Name.Builtin: RGBColor(0x00, 0x00, 0xFF),
                        Token.Name.Function: RGBColor(0x00, 0x00, 0x00),
                        Token.String: RGBColor(0xA3, 0x15, 0x15),
                        Token.Comment: RGBColor(0x00, 0x80, 0x00),
                        Token.Comment.Preproc: RGBColor(0x00, 0x00, 0xFF),
                        Token.Number: RGBColor(0x09, 0x80, 0x85),
                        Token.Operator: RGBColor(0x00, 0x00, 0x00),
                        Token.Name.Class: RGBColor(0x2B, 0x91, 0xAF),
                        Token.Generic: RGBColor(0x00, 0x00, 0x00),
                    }
                    is_keyword = token_type in Token.Keyword or token_type in Token.Keyword.Type
                    is_comment = token_type in Token.Comment or token_type in Token.Comment.Preproc

                    color = None
                    for tk, c in color_map.items():
                        if token_type in tk or tk in token_type:
                            color = c
                            break

                    add_code_run(code_para, token_text, color=color,
                               bold=is_keyword, italic=is_comment)

        # 题间分隔
        sep_para = doc.add_paragraph()
        sep_para.paragraph_format.space_before = Pt(8)
        pPr = sep_para._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="4" w:space="1" w:color="E0E0E0"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

    # 章尾
    end_para = doc.add_paragraph()
    end_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end_para.paragraph_format.space_before = Pt(24)
    end_run = end_para.add_run('— 第1章完 · 共14题 —')
    end_run.font.size = Pt(8)
    end_run.font.color.rgb = LIGHT_GRAY

    out = BOOK_ROOT / "textbook" / "chapter01.docx"
    doc.save(str(out))
    return out


# ============================================================
if __name__ == "__main__":
    html_path = generate_html()
    print(f"✅ HTML: {html_path}")

    docx_path = generate_docx()
    print(f"✅ DOCX: {docx_path}")
