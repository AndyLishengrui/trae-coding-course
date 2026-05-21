import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from pygments import lex
from pygments.lexers import CppLexer, PythonLexer
from pygments.styles import get_style_by_name
from pygments.token import Token

# 配置
ROOT_DIR = "NQ100"
FONT_NAME = "Consolas" # 建议字体
FONT_SIZE = 8.5        # 建议字号
STYLE_NAME = "default" # Pygments 样式

def hex_to_rgb(hex_str):
    """将十六进制颜色转换为 (r, g, b) 元组"""
    if not hex_str:
        return None
    hex_str = hex_str.lstrip('#')
    return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

def set_cell_right_border(cell):
    """
    给单元格添加右边框，用于实现中间分割线。
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    # 定义右边框属性
    # w:val="single" 单实线
    # w:sz="6" (1/8 pt 单位) -> 0.75pt
    # w:color="auto" 自动颜色(黑)
    right = OxmlElement('w:right')
    right.set(qn('w:val'), 'single')
    right.set(qn('w:sz'), '6')
    right.set(qn('w:space'), '0')
    right.set(qn('w:color'), 'auto')
    
    # 如果已有右边框定义则替换，否则追加
    existing = tcBorders.find(qn('w:right'))
    if existing is not None:
        tcBorders.remove(existing)
    tcBorders.append(right)

def add_highlighted_code(cell, code_text, lexer, style):
    """
    将代码高亮添加到表格单元格中
    """
    # 添加新段落用于存放代码
    paragraph = cell.add_paragraph()
    
    # 设置段落格式
    paragraph_format = paragraph.paragraph_format
    paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0
    paragraph_format.left_indent = Pt(0)
    paragraph_format.first_line_indent = Pt(0)
    
    # 解析代码
    tokens = lex(code_text.strip(), lexer)
    
    for token_type, value in tokens:
        # 获取样式
        token_style = style.style_for_token(token_type)
        
        run = paragraph.add_run(value)
        font = run.font
        font.name = FONT_NAME
        font.size = Pt(FONT_SIZE)
        
        # 颜色
        if token_style['color']:
            rgb = hex_to_rgb(token_style['color'])
            if rgb:
                font.color.rgb = RGBColor(*rgb)
        
        # 粗体/斜体
        if token_style['bold']:
            font.bold = True
        # 强制不使用斜体，因为在 Word 中显示效果不佳（特别是注释）
        if token_style['italic']:
            font.italic = False

def parse_markdown_to_docx(document, lines):
    """
    简单解析 Markdown 并添加到 Word 文档
    支持: 
    - #, ##, ### (Heading)
    - 1. (List Number)
    - -, * (List Bullet)
    - **Bold** (Bold Text)
    """
    for line in lines:
        content = line.strip()
        if not content:
            continue

        style = None
        
        # Headings
        if content.startswith('# '):
            document.add_heading(content[2:], level=1)
            continue
        elif content.startswith('## '):
            document.add_heading(content[3:], level=2)
            continue
        elif content.startswith('### '):
            document.add_heading(content[4:], level=3)
            continue
            
        # Lists
        match_ordered = re.match(r'^(\d+\.)\s+(.*)', content)
        if match_ordered:
            style = 'List Number'
            content = match_ordered.group(2)
        elif content.startswith('- ') or content.startswith('* '):
            style = 'List Bullet'
            content = content[2:]
        
        # Add Paragraph
        if style:
            try:
                p = document.add_paragraph(style=style)
            except:
                p = document.add_paragraph()
        else:
            p = document.add_paragraph()
            
        # Process Bold markers **text**
        parts = re.split(r'\*\*(.*?)\*\*', content)
        for i, part in enumerate(parts):
            if not part: continue
            run = p.add_run(part)
            if i % 2 == 1:
                run.bold = True

def process_problem(problem_num):
    folder_name = f"NQ{int(problem_num):03d}"
    full_path = os.path.join(ROOT_DIR, folder_name)
    
    if not os.path.exists(full_path):
        print(f"Error: 找不到题目目录 {folder_name} (路径: {full_path})")
        return

    output_file = os.path.join(full_path, f"{folder_name}.docx")
    print(f"正在生成 {output_file} ...")
    
    document = Document()
    
    # 设置页面边距，尽量宽一点以便放下双栏
    sections = document.sections
    for section in sections:
        section.left_margin = Cm(1.27)   # 0.5 inch
        section.right_margin = Cm(1.27)
        section.page_width = Cm(21.0)    # A4 width
        section.page_height = Cm(29.7)   # A4 height

    style = get_style_by_name(STYLE_NAME)

    # 1. 查找思路文件
    thinking_file = os.path.join(full_path, "思路.md")
    
    # 2. 查找 C++ 文件
    # 优先级：NQXXX_improved.cpp > cpp_compressed.cpp > NQXXX.cpp > std.cpp
    cpp_file = None
    candidates_cpp = [
        os.path.join(full_path, f"{folder_name}_improved.cpp"),
        os.path.join(full_path, "cpp_compressed.cpp"),
        os.path.join(full_path, f"{folder_name}.cpp"),
        os.path.join(full_path, "std.cpp")
    ]
    for c in candidates_cpp:
        if os.path.exists(c):
            cpp_file = c
            break
    
    if not cpp_file:
        # 稍微放宽文件查找条件，找任意cpp
        possible_cpp = [f for f in os.listdir(full_path) if f.endswith(".cpp") and "std" not in f and "compressed" not in f]
        if possible_cpp:
            cpp_file = os.path.join(full_path, possible_cpp[0])
    
    # 3. 查找 Python 文件
    # 优先级：NQXXX_improved.py > solution.py
    py_file = None
    candidates_py = [
        os.path.join(full_path, f"{folder_name}_improved.py"),
        os.path.join(full_path, "solution.py")
    ]
    for p in candidates_py:
        if os.path.exists(p):
            py_file = p
            break
            
    if not py_file:
        possible_py = [f for f in os.listdir(full_path) if f.endswith(".py") and "solution" not in f]
        if possible_py:
            py_file = os.path.join(full_path, possible_py[0])

    has_cpp = cpp_file and os.path.exists(cpp_file)
    has_py = py_file and os.path.exists(py_file)

    if has_cpp or has_py:
        # --- 第一部分：解题思路 ---
        # 优先读取并解析思路文件，如果不存在或者没有标题，则使用默认标题
        
        has_h1 = False
        thinking_lines = []
        if os.path.exists(thinking_file):
            try:
                with open(thinking_file, 'r', encoding='utf-8') as f:
                    thinking_lines = f.readlines()
                # 检查是否包含H1
                for line in thinking_lines:
                    if line.strip().startswith('# '):
                        has_h1 = True
                        break
            except Exception as e:
                print(f"读取思路文件失败: {e}")

        # 如果Markdown中没有一级标题，则添加默认标题
        if not has_h1:
            document.add_heading(folder_name, level=1)
        
        if thinking_lines:
            parse_markdown_to_docx(document, thinking_lines)
        else:
            document.add_heading("解题思路", level=2)
            document.add_paragraph("暂无解题思路。")

        # --- 第二部分：参考代码 ---
        document.add_heading("参考代码", level=2)
        
        # 创建 1 行 2 列的表格
        table = document.add_table(rows=1, cols=2)
        # table.style = 'Table Grid' # 移除全边框
        table.autofit = False 
        
        # 设置列宽（均匀分布）
        col_width = (document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin) / 2
        table.columns[0].width = int(col_width)
        table.columns[1].width = int(col_width)
        
        # --- C++ Column (Left) ---
        cell_left = table.rows[0].cells[0]
        set_cell_right_border(cell_left) # 设置中间分割线
        
        p_left_header = cell_left.paragraphs[0]
        r_left = p_left_header.add_run("C++ Code")
        r_left.bold = True
        r_left.font.color.rgb = RGBColor(0, 0, 139) # DarkBlue
        p_left_header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        if has_cpp:
            try:
                with open(cpp_file, 'r', encoding='utf-8') as f:
                    code_content = f.read()
                add_highlighted_code(cell_left, code_content, CppLexer(), style)
            except Exception as e:
                cell_left.add_paragraph(f"Error: {e}")
        else:
            cell_left.add_paragraph("No C++ file.")

        # --- Python Column (Right) ---
        cell_right = table.rows[0].cells[1]
        p_right_header = cell_right.paragraphs[0]
        r_right = p_right_header.add_run("Python Code")
        r_right.bold = True
        r_right.font.color.rgb = RGBColor(0, 100, 0) # DarkGreen
        p_right_header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        if has_py:
            try:
                with open(py_file, 'r', encoding='utf-8') as f:
                    code_content = f.read()
                add_highlighted_code(cell_right, code_content, PythonLexer(), style)
            except Exception as e:
                    cell_right.add_paragraph(f"Error: {e}")
        else:
            cell_right.add_paragraph("No Python file.")

        document.save(output_file)
        print(f"成功生成文档: {output_file}")
    else:
        print(f"在 {folder_name} 中未找到 C++ 或 Python 代码文件，跳过生成。")

def main():
    if len(sys.argv) > 1:
        problem_num = sys.argv[1]
    else:
        problem_num = input("请输入题目编号 (例如 1): ")
    
    try:
        process_problem(problem_num)
    except ValueError:
        print("请输入有效的数字编号")
    except Exception as e:
        print(f"发生错误: {e}")

if __name__ == "__main__":
    main()
