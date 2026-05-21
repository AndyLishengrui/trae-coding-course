import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pygments import lex
from pygments.lexers import CppLexer, PythonLexer
from pygments.styles import get_style_by_name
from pygments.token import Token

# 配置
ROOT_DIR = "NQ100"
OUTPUT_FILE = "NQ100_Code_Collection.docx"
FONT_NAME = "Courier New" # 等宽字体
FONT_SIZE = 9
STYLE_NAME = "default" # Pygments 样式名，可选 'friendly', 'colorful', 'monokai' 等

def hex_to_rgb(hex_str):
    """将十六进制颜色转换为 (r, g, b) 元组"""
    if not hex_str:
        return None
    hex_str = hex_str.lstrip('#')
    return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

def add_highlighted_code(paragraph, code_text, lexer, style):
    """
    将代码高亮添加到 docx 的段落中
    """
    # 设置段落格式
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0
    
    # 模拟背景色（Word段落背景色设置比较复杂，这里主要关注文字高亮）
    # 如果需要背景色，建议放在表格单元格中设置
    
    # 解析代码
    tokens = lex(code_text, lexer)
    
    for token_type, value in tokens:
        # 获取样式
        token_style = style.style_for_token(token_type)
        
        run = paragraph.add_run(value)
        font = run.font
        font.name = FONT_NAME
        font.size = Pt(FONT_SIZE)
        
        # 颜色
        if token_style['color']:
            rgb = hex_to_rgb(token_style['color'])
            if rgb:
                font.color.rgb = RGBColor(*rgb)
        
        # 粗体/斜体
        if token_style['bold']:
            font.bold = True
        if token_style['italic']:
            font.italic = True

def main():
    document = Document()
    
    # 设置页面边距（可选）
    sections = document.sections
    for section in sections:
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    style = get_style_by_name(STYLE_NAME)
    
    # 获取题目列表并排序
    valid_dirs = []
    if os.path.exists(ROOT_DIR):
        for d in os.listdir(ROOT_DIR):
            if d.startswith("NQ") and os.path.isdir(os.path.join(ROOT_DIR, d)):
                valid_dirs.append(d)
    
    # 自定义排序：NQ1, NQ2, ..., NQ100
    valid_dirs.sort(key=lambda x: int(re.search(r'\d+', x).group()) if re.search(r'\d+', x) else 0)

    for problem_dir in valid_dirs:
        full_path = os.path.join(ROOT_DIR, problem_dir)
        print(f"Processing {problem_dir}...")
        
        # 优先使用压缩版代码
        cpp_file = os.path.join(full_path, "cpp_compressed.cpp")
        if not os.path.exists(cpp_file):
            cpp_file = os.path.join(full_path, "std.cpp")
        
        py_file = os.path.join(full_path, "solution.py")
        
        # 稍微放宽文件查找条件，有时可能叫 NQxxx.cpp
        if not os.path.exists(cpp_file):
            possible_cpp = [f for f in os.listdir(full_path) if f.endswith(".cpp") and "std" not in f and "compressed" not in f]
            if possible_cpp:
                cpp_file = os.path.join(full_path, possible_cpp[0])
        
        if not os.path.exists(py_file):
             possible_py = [f for f in os.listdir(full_path) if f.endswith(".py") and "solution" not in f]
             if possible_py:
                py_file = os.path.join(full_path, possible_py[0])

        # 只要有一个文件存在就处理
        if os.path.exists(cpp_file) or os.path.exists(py_file):
            # 添加题目大标题
            document.add_heading(problem_dir, level=1)
            
            # 创建一个 2行1列 的表格来实现“上下栏”效果
            table = document.add_table(rows=2, cols=1)
            table.style = 'Table Grid' #甚至是 'Normal Table' 隐藏边框
            
            # --- C++ 部分 ---
            cell_top = table.rows[0].cells[0]
            # 添加标签
            p_label = cell_top.add_paragraph()
            run = p_label.add_run("C++ Code")
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 139) # DarkBlue
            
            if os.path.exists(cpp_file):
                try:
                    with open(cpp_file, 'r', encoding='utf-8') as f:
                        code_content = f.read()
                    p_code = cell_top.add_paragraph()
                    add_highlighted_code(p_code, code_content, CppLexer(), style)
                except Exception as e:
                    cell_top.add_paragraph(f"Error reading C++ file: {e}")
            else:
                cell_top.add_paragraph("No C++ file found.")

            # --- Python 部分 ---
            cell_bottom = table.rows[1].cells[0]
            # 添加标签
            p_label = cell_bottom.add_paragraph()
            run = p_label.add_run("Python Code")
            run.bold = True
            run.font.color.rgb = RGBColor(0, 100, 0) # DarkGreen
            
            if os.path.exists(py_file):
                try:
                    with open(py_file, 'r', encoding='utf-8') as f:
                        code_content = f.read()
                    p_code = cell_bottom.add_paragraph()
                    add_highlighted_code(p_code, code_content, PythonLexer(), style)
                except Exception as e:
                     cell_bottom.add_paragraph(f"Error reading Python file: {e}")
            else:
                cell_bottom.add_paragraph("No Python file found.")
            
            # 在每题之间添加分页符
            document.add_page_break()

    document.save(OUTPUT_FILE)
    print(f"Document generated: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
